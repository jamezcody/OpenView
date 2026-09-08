"""Additional regulator exports, with source-specific location semantics.

Schemas verified against official exports in September 2026. The Dutch export
contains monthly worksheets; only the newest populated month is emitted.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
import zipfile
import xml.etree.ElementTree as ET
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse

from openpyxl import load_workbook
from core import NeedsImport, dms, norm, number

NL_PAGE = 'https://www.antennebureau.nl/plaatsing-antennes/aantal-antenne-installaties-voor-mobiele-communicatie'
PL_PAGE = 'https://bip.uke.gov.pl/pozwolenia-radiowe/wykaz-pozwolen-radiowych-tresci/stacje-gsm-umts-lte-5gnr-oraz-cdma,12,91.html'
CL_PAGE = 'https://www.subtel.gob.cl/ubicacion-antenas-5g/'
SOURCES = {
    'nl_rdi': dict(name='Netherlands — RDI cellular antennas', country='NL', mode='automatic',
        url=NL_PAGE, native_extensions=['.ods'],
        scope='Official cellular-antenna ODS export; latest populated monthly worksheet only. WGS84 coordinates have a stated 15 m accuracy. No full Antenneregister map scraping.'),
    'de_bnetza': dict(name='Germany — BNetzA EMF sites', country='DE', mode='import',
        url='https://www.bundesnetzagentur.de/DE/Fachthemen/Telekommunikation/Technik/EMF/start.html',
        scope='Mapped CSV import only. Obtain a permitted site-location table from BNetzA or your authorized EMF portal access; map latitude, longitude, certificate/site ID and any available technical columns. Public EMF map has no verified unrestricted bulk export; detailed portal is restricted to competent authorities. Certificates can include planned installations. Do not import EMF measurement points as transmitters.'),
    'pl_uke': dict(name='Poland — UKE cellular permits', country='PL', mode='automatic',
        url=PL_PAGE, native_extensions=['.xlsx', '.zip'],
        scope='All XLSX permit lists linked on the current UKE cellular page, including private networks. Authorization does not establish that a base station was built or operates. Nominal band labels are not treated as exact transmitting frequencies.'),
    'cl_subtel': dict(name='Chile — SUBTEL published antenna locations', country='CL', mode='automatic',
        url=CL_PAGE, native_extensions=['.xlsx'],
        scope='Published 2021 selected-band location workbook (700 MHz, AWS, 3.5 GHz and 26 GHz worksheets). Static publication; no assertion of current deployment or operation. Preserve worksheet band labels without inventing exact channel frequencies.'),
}


class _Links(HTMLParser):
    def __init__(self):
        super().__init__(); self.links = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            link = dict(attrs).get('href')
            if link: self.links.append(link)


def _links(ctx, url):
    page = ctx.download(url, 'source-index.html')
    parser = _Links(); parser.feed(page.read_text(encoding='utf-8-sig'))
    # Automatic downloads stay on the verified official publisher's host.
    host = urlparse(url).hostname
    return list(dict.fromkeys(urljoin(url, link) for link in parser.links
        if urlparse(urljoin(url, link)).scheme == 'https'
        and urlparse(urljoin(url, link)).hostname == host))


def acquire(ctx, source):
    if source == 'de_bnetza':
        raise NeedsImport(SOURCES[source]['scope'] + ' Open ' + SOURCES[source]['url'])
    if source == 'nl_rdi':
        links = _links(ctx, NL_PAGE)
        documents = [u for u in links if re.search(r'/documenten/\d{4}/', u)
                     and any(s in u.lower() for s in ('jaaroverzicht', 'overzicht-aantal'))]
        if not documents: raise ValueError('RDI page has no supported annual export link')
        latest = max(documents, key=lambda u: re.search(r'/documenten/(\d{4})/', u)[1])
        files = [u for u in _links(ctx, latest) if urlparse(u).path.lower().endswith('.ods')]
        if len(files) != 1: raise ValueError('Expected one ODS export on latest RDI annual page')
    elif source in ('pl_uke', 'cl_subtel'):
        files = [u for u in _links(ctx, SOURCES[source]['url'])
                 if urlparse(u).path.lower().endswith('.xlsx')]
        if not files: raise ValueError('No XLSX exports found on ' + SOURCES[source]['url'])
    else:
        raise ValueError('Unknown national source: ' + source)
    for url in files:
        yield ctx.download(url, Path(unquote(urlparse(url).path)).name), url


def _base(source, identity, filename, url):
    return dict(source=source, country=SOURCES[source]['country'],
        source_record_id=identity, source_file=filename, source_url=url,
        source_crs='unspecified', location_kind='fixed', transmission_kind='transmit')


def _fingerprint(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, default=str,
                                    ensure_ascii=False).encode('utf-8')).hexdigest()


def _iso(value):
    if isinstance(value, (datetime, date)): return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return str(value or '').strip()


def coordinate(value, hemisphere):
    """Official DMS variants, packed DDMMSS, or already signed decimal degrees."""
    text = str(value if value is not None else '').strip().upper().replace(',', '.')
    if not text: return None
    explicit = re.findall(r'[NSEW]', text)
    direction = explicit[0] if len(explicit) == 1 else hemisphere
    if len(explicit) > 1 or direction not in ('N', 'S', 'E', 'W'): return None
    if (direction in 'NS') != (hemisphere in 'NS'): return None
    pieces = re.findall(r'\d+(?:\.\d+)?', text)
    if len(pieces) == 3:
        result = dms(*pieces, direction)
    elif len(pieces) == 1:
        value = number(pieces[0])
        if value is None: return None
        if value > 180:
            if '.' in pieces[0] or len(pieces[0]) not in (5, 6, 7): return None
            result = dms(pieces[0][:-4], pieces[0][-4:-2], pieces[0][-2:], direction)
        else:
            result = value * (-1 if direction in 'SW' else 1)
    else: return None
    if result is not None and text.startswith('-'): result = -abs(result)
    limit = 90 if hemisphere in 'NS' else 180
    return result if result is not None and abs(result) <= limit else None


def _workbooks(path):
    path = Path(path)
    if path.suffix.lower() != '.zip':
        book = load_workbook(path, read_only=True, data_only=True)
        try: yield path.name, book
        finally: book.close()
        return
    with zipfile.ZipFile(path) as archive:
        files = [n for n in archive.namelist() if n.lower().endswith('.xlsx')]
        if not files: raise ValueError('Permit ZIP contains no XLSX files')
        for member in files:
            # Read in memory, never extract paths supplied by the archive.
            book = load_workbook(io.BytesIO(archive.read(member)), read_only=True, data_only=True)
            try: yield member, book
            finally: book.close()


def parse_poland(ctx, ds, path, url):
    required = {norm(s) for s in ('Nazwa Operatora', 'Nr Decyzji', 'Dł geogr stacji', 'Szer geogr stacji')}
    for filename, book in _workbooks(path):
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            headers = None
            for line, values in enumerate(rows, 1):
                if line % 5000 == 0: ctx.check()
                if not any(v is not None and str(v).strip() for v in values): continue
                if headers is None:
                    keys = [norm(v) for v in values]
                    if required <= set(keys): headers = keys; continue
                    if line >= 10: raise ValueError('Unsupported UKE worksheet headers: ' + sheet.title)
                    continue
                r = dict(zip(headers, values))
                if not r.get(norm('Nr Decyzji')): continue
                band = re.match(r'(5g|lte|gsm|umts|cdma)\s*(\d+)', sheet.title.strip().lower())
                band_label = band[0] if band else sheet.title
                expiry = _iso(r.get(norm('Data ważności')))
                out = _base('pl_uke', band_label + ':' + _fingerprint(r), filename + '#' + sheet.title, url)
                out.update(operator=r.get(norm('Nazwa Operatora')), license_id=r.get(norm('Nr Decyzji')),
                    site_id=r.get('idstacji'), latitude=coordinate(r.get(norm('Szer geogr stacji')), 'N'),
                    longitude=coordinate(r.get(norm('Dł geogr stacji')), 'E'), service='cellular ' + band_label,
                    expiration_date=expiry, status='expired' if re.fullmatch(r'\d{4}-\d{2}-\d{2}', expiry)
                    and expiry < date.today().isoformat() else 'authorized_not_operationally_verified',
                    details_json={'worksheet': sheet.title, 'decision_type': r.get(norm('Rodzaj decyzji')),
                        'locality': r.get(norm('Miejscowość')), 'address': r.get('lokalizacja'),
                        'nominal_band_mhz': int(band[2]) if band else None,
                        'authorization_is_not_evidence_of_operation': True, 'source_row': line})
                frequency = str(r.get(norm('Zakres częstotliwości [MHz]')) or '')
                match = re.fullmatch(r'\s*(\d+(?:[.,]\d+)?)\s*-\s*(\d+(?:[.,]\d+)?)\s*', frequency)
                if match:
                    out.update(frequency_mhz=number(match[1]), frequency_upper_mhz=number(match[2]))
                ds.add(out)
            if headers is None: raise ValueError('No supported UKE header in ' + sheet.title)


def parse_chile(ctx, ds, path, url):
    for filename, book in _workbooks(path):
        for sheet in book:
            columns = {}; buffer = []
            for line, values in enumerate(sheet.iter_rows(values_only=True), 1):
                if line % 5000 == 0: ctx.check()
                if len(buffer) < 10 and not {'lat', 'lon', 'name'} <= set(columns):
                    buffer.append(values)
                    for i, value in enumerate(values):
                        key = norm(value)
                        if key.startswith('latitud'): columns['lat'] = i
                        elif key.startswith('longitud'): columns['lon'] = i
                        elif key.startswith('nombre'): columns['name'] = i
                        elif key.startswith('direccion'): columns['address'] = i
                        elif key == 'tecnologia': columns['technology'] = i
                    if line >= 10 and not {'lat', 'lon', 'name'} <= set(columns):
                        raise ValueError('Unsupported SUBTEL worksheet headers: ' + sheet.title)
                    continue
                def get(key):
                    i = columns.get(key)
                    return values[i] if i is not None and i < len(values) else None
                if not get('lat') and not get('lon'): continue
                lat = coordinate(get('lat'), 'S'); lon = coordinate(get('lon'), 'W')
                out = _base('cl_subtel', sheet.title + ':' + _fingerprint(values), filename + '#' + sheet.title, url)
                # The sheet and country convention supplies S/W where omitted.
                # WGS84 is stated in the 700/AWS/26 headers; 3500 headers omit datum.
                datum = 'EPSG:4326' if any('WGS84' in str(v) for row in buffer for v in row) else 'unspecified'
                out.update(site_id=get('name'), operator=sheet.title.rsplit(' ', 1)[0],
                    latitude=lat, longitude=lon, source_crs=datum,
                    service=str(get('technology') or 'cellular') + ' / ' + sheet.title,
                    status='published_location_not_operationally_verified',
                    details_json={'worksheet': sheet.title, 'source_row': line, 'address': get('address'),
                        'publication_date': '2021-03-27', 'band_label': sheet.title.rsplit(' ', 1)[-1],
                        'coordinate_hemisphere': 'S/W, explicit when present; otherwise Chile convention',
                        'current_operation_unverified': True})
                ds.add(out)
            if not {'lat', 'lon', 'name'} <= set(columns):
                raise ValueError('No supported SUBTEL header in ' + sheet.title)


_TABLE = '{urn:oasis:names:tc:opendocument:xmlns:table:1.0}'
_OFFICE = '{urn:oasis:names:tc:opendocument:xmlns:office:1.0}'
_MONTHS = {v: i for i, v in enumerate(('jan','feb','mrt','apr','mei','jun','jul','aug','sep','okt','nov','dec'), 1)}
_MONTHS['maa'] = 3


def ods_rows(path):
    """Bounded column expansion avoids ODS's 16,384-column blank repetitions."""
    with zipfile.ZipFile(path) as archive, archive.open('content.xml') as stream:
        table_name = ''; table_element = None
        for event, element in ET.iterparse(stream, events=('start', 'end')):
            if event == 'start' and element.tag == _TABLE + 'table':
                table_name = element.get(_TABLE + 'name', ''); table_element = element
            elif event == 'end' and element.tag == _TABLE + 'table-row':
                values = []
                for cell in element:
                    if cell.tag not in (_TABLE + 'table-cell', _TABLE + 'covered-table-cell'): continue
                    value = cell.get(_OFFICE + 'value', ''.join(cell.itertext())).strip()
                    repeat = int(cell.get(_TABLE + 'number-columns-repeated', '1'))
                    values.extend([value] * min(repeat, max(0, 32 - len(values))))
                    if len(values) >= 32: break
                if any(values):
                    repeated = int(element.get(_TABLE + 'number-rows-repeated', '1'))
                    if repeated > 100000: raise ValueError('Unreasonable repeated populated ODS rows')
                    for _ in range(repeated): yield table_name, values
                element.clear()
                if table_element is not None: table_element.clear()


def parse_netherlands(ctx, ds, path, url, stage_path):
    db = sqlite3.connect(stage_path)
    db.execute('DROP TABLE IF EXISTS nl_latest')
    db.execute('CREATE TABLE nl_latest (record TEXT)')
    latest = 0; headers = {}; input_rows = 0
    required = {'gemeente', 'toepassing', 'coordinatennoorderbreedte', 'coordinatenoosterlengte', 'antennehoogte'}
    try:
        for sheet, values in ods_rows(path):
            input_rows += 1
            if input_rows % 5000 == 0: ctx.check()
            keys = [norm(v) for v in values]
            if required <= set(keys): headers[sheet] = keys; continue
            if sheet not in headers: continue
            r = dict(zip(headers[sheet], values))
            lat = number(r.get('coordinatennoorderbreedte')); lon = number(r.get('coordinatenoosterlengte'))
            if lat is None and lon is None: continue
            month = _MONTHS.get(norm(sheet)[:3])
            if not month: raise ValueError('Unknown Dutch monthly worksheet name: ' + sheet)
            if month < latest: continue
            if month > latest: db.execute('DELETE FROM nl_latest'); latest = month
            out = _base('nl_rdi', _fingerprint(r), Path(path).name + '#' + sheet, url)
            out.update(latitude=lat, longitude=lon, source_crs='EPSG:4326',
                service='cellular ' + str(r.get('toepassing', '')), antenna_height_m=r.get('antennehoogte'),
                status='registered_installation_as_published',
                details_json={'worksheet': sheet, 'month': month, 'municipality': r.get('gemeente'),
                    'stated_coordinate_accuracy_m': 15, 'location_record_not_individual_radio': True})
            db.execute('INSERT INTO nl_latest VALUES (?)', (json.dumps(out, ensure_ascii=False),))
        if not latest: raise ValueError('No populated location worksheet in Dutch ODS export')
        ctx.log('Using Dutch monthly worksheet ' + str(latest) + ' from published annual export')
        for row, in db.execute('SELECT record FROM nl_latest'):
            ds.add(json.loads(row))
    finally:
        db.close()


def parse(ctx, ds, source, path, url, stage_path):
    if source == 'pl_uke': parse_poland(ctx, ds, path, url)
    elif source == 'cl_subtel': parse_chile(ctx, ds, path, url)
    elif source == 'nl_rdi': parse_netherlands(ctx, ds, path, url, stage_path)
    elif source == 'de_bnetza': raise NeedsImport('Use a CSV mapping for the authorized BNetzA site table')
    else: raise ValueError('Unknown national source: ' + source)
